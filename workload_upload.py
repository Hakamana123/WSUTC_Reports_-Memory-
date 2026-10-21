"""Bulk upload: the staff list, and optionally their teaching, from one Excel file.

The template has a "Staff" sheet (one row per person) and a "Teaching" sheet
(one line per subject a person teaches in a session / block), with dropdowns
for the closed lists. Headers are found by name, not position, so columns can
be moved or left out — except Name, which every row needs.

Staff are matched to the list by Employee No, then by name (as the load
tracking import does). A match is updated from the cells that are filled in —
a blank cell keeps what's there — so the same file can be uploaded again after
edits. A teaching line already on file for the same person, session, block
and subject is updated rather than added twice.

Rows with a problem are skipped and listed; nothing is saved here — the page
shows the preview, adds the rows and saves.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import workload_config as cfg
import workload_store as wstore
from workload_import import _emp, _num, _text, name_key

STAFF_SHEET, TEACHING_SHEET = "Staff", "Teaching"

# template column -> (stored column, header prefixes accepted on upload)
STAFF_COLS = {
    "Name": ("Staff", ["name", "staff"]),
    "Employee No": ("Employee No", ["employee"]),
    "Role": ("Role", ["role"]),
    "FTE": ("FTE", ["fte"]),
    "Supervisor": ("Supervisor", ["supervisor"]),
    "Home program": ("Home discipline", ["home"]),
    "Active": ("Active", ["active"]),
    "Notes": ("Notes", ["note"]),
}
TEACHING_COLS = {
    "Name": ("Staff", ["name", "staff"]),
    "Session": ("Session", ["session"]),
    "Block": ("Block", ["block"]),
    "Program": ("Discipline", ["program", "discipline"]),
    "Subject": ("Subject", ["subject", "unit"]),
    "Classes": ("Classes", ["class"]),
    "DI hrs/wk": ("DI hrs/wk", ["di", "hrs", "hours"]),
    "Notes": ("Notes", ["note"]),
}
YES, NO = "TRUE", "FALSE"
SESSION_RE = re.compile(r"^(\d{2}) (AUT|SPR|SUM)$")


# --------------------------------------------------------------------------
# the template
# --------------------------------------------------------------------------


def template_xlsx(programs: list[str], sessions: list[str]) -> bytes:
    """An empty upload file: Staff and Teaching sheets with dropdowns, plus a
    hidden Lists sheet the dropdowns read from."""
    wb = Workbook()
    lists = {
        "Role": cfg.STAFF_ROLES,
        "Program": programs,
        "Active": [YES, NO],
        "Session": sessions,
        "Block": cfg.BLOCK_OPTIONS,
    }
    ws_lists = wb.active
    ws_lists.title = "Lists"
    refs = {}
    for j, (name, values) in enumerate(lists.items(), 1):
        col = get_column_letter(j)
        ws_lists.cell(row=1, column=j, value=name)
        for i, v in enumerate(values, 2):
            ws_lists.cell(row=i, column=j, value=v)
        refs[name] = f"Lists!${col}$2:${col}${max(len(values), 1) + 1}"
    ws_lists.sheet_state = "hidden"

    def sheet(title, cols, dropdowns, example):
        ws = wb.create_sheet(title)
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=c)
            cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="EDEDED")
            ws.column_dimensions[get_column_letter(j)].width = 28 if c in ("Name", "Program", "Home program",
                                                                          "Subject", "Notes", "Role") else 13
            if c in dropdowns:
                dv = DataValidation(type="list", formula1=f"={refs[dropdowns[c]]}", allow_blank=True)
                dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}1000")
                ws.add_data_validation(dv)
        for j, v in enumerate(example, 1):
            ws.cell(row=2, column=j, value=v).font = Font(italic=True, color="888888")
        ws.freeze_panes = "A2"

    first = sessions[0] if sessions else ""
    prog = programs[0] if programs else ""
    sheet(STAFF_SHEET, list(STAFF_COLS),
          {"Role": "Role", "Home program": "Program", "Active": "Active"},
          ["Example, Pat (delete this row)", "12345678", cfg.DEFAULT_ROLE, 1, "Supervisor name", prog, YES, ""])
    sheet(TEACHING_SHEET, list(TEACHING_COLS),
          {"Session": "Session", "Block": "Block", "Program": "Program"},
          ["Example, Pat (delete this row)", first, cfg.ALL_BLOCKS, prog, "ABCD1001", 1, 4, ""])
    wb.move_sheet("Lists", offset=2)
    wb.active = 0
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------
# reading an upload
# --------------------------------------------------------------------------


@dataclass
class Upload:
    staff: list[dict] = field(default_factory=list)      # stored column -> value, plus "_row"
    teaching: list[dict] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)    # rows skipped
    notes: list[str] = field(default_factory=list)       # kept, but worth a look


def _read_sheet(sheets: dict, wanted: str, cols: dict) -> tuple[list[dict], str | None]:
    """Rows of the sheet named `wanted` (case-insensitive) as {stored column:
    text}, plus "_row" (the Excel row number). Example rows are dropped."""
    raw = next((df for name, df in sheets.items() if name.strip().lower() == wanted.lower()), None)
    if raw is None or raw.empty:
        return [], None
    header = [_text(h).lower() for h in raw.iloc[0]]
    # a header starting with the word first ("Name"), then one with it as a
    # later word ("Full name") for whatever is still missing
    where = {}
    for test in (lambda h, p: h.startswith(p), lambda h, p: f" {p}" in f" {h}"):
        for _, (stored, prefixes) in cols.items():
            for j, h in enumerate(header):
                if stored not in where and j not in where.values() and any(test(h, p) for p in prefixes):
                    where[stored] = j
                    break
    if "Staff" not in where:
        return [], f"The '{wanted}' sheet has no Name column — skipped."
    rows = []
    for i in range(1, len(raw)):
        r = {stored: _text(raw.iat[i, j]) for stored, j in where.items()}
        if not any(r.values()) or "(delete this row)" in r["Staff"]:
            continue
        r["_row"] = i + 1
        rows.append(r)
    return rows, None


def _pick(value: str, options: list[str]) -> str | None:
    """`value` matched to one of `options` ignoring case; None if no match."""
    return next((o for o in options if o.lower() == value.strip().lower()), None)


def _number(v: str) -> str:
    n = _num(v)
    return "" if n is None else (str(int(n)) if float(n).is_integer() else str(round(n, 2)))


def parse_upload(file, programs: list[str]) -> Upload:
    """Read and check an upload file. `programs` = the known programs."""
    sheets = pd.read_excel(file, sheet_name=None, header=None, dtype=object, keep_default_na=False)
    up = Upload()

    staff_rows, err = _read_sheet(sheets, STAFF_SHEET, STAFF_COLS)
    if err:
        up.problems.append(err)
    for r in staff_rows:
        where = f"Staff row {r['_row']}"
        if not r["Staff"]:
            up.problems.append(f"{where}: no name — skipped.")
            continue
        if r.get("Role"):
            role = _pick(r["Role"], cfg.STAFF_ROLES)
            if role is None:
                up.problems.append(f"{where} ({r['Staff']}): unknown role '{r['Role']}' — skipped. "
                                   f"Use one of: {', '.join(cfg.STAFF_ROLES)}.")
                continue
            r["Role"] = role
        if r.get("FTE"):
            f = _num(r["FTE"])
            if f is None or not 0 < f <= 1:
                up.problems.append(f"{where} ({r['Staff']}): FTE '{r['FTE']}' isn't between 0 and 1 — skipped.")
                continue
            r["FTE"] = _number(r["FTE"])
        if r.get("Active"):
            a = r["Active"].strip().lower()
            r["Active"] = NO if a in ("false", "no", "n", "0", "inactive") else YES
        if r.get("Employee No"):
            r["Employee No"] = _emp(r["Employee No"])
        if r.get("Home discipline"):
            known = _pick(r["Home discipline"], programs)
            if known is None:
                up.notes.append(f"{where} ({r['Staff']}): program '{r['Home discipline']}' isn't on the "
                                "list — kept as typed.")
            r["Home discipline"] = known or r["Home discipline"]
        up.staff.append(r)

    teach_rows, err = _read_sheet(sheets, TEACHING_SHEET, TEACHING_COLS)
    if err:
        up.problems.append(err)
    for r in teach_rows:
        where = f"Teaching row {r['_row']}"
        who = f" ({r['Staff']})" if r["Staff"] else ""
        if not r["Staff"]:
            up.problems.append(f"{where}: no name — skipped.")
            continue
        m = SESSION_RE.match(r.get("Session", "").strip().upper())
        if not m:
            up.problems.append(f"{where}{who}: session '{r.get('Session', '')}' should look like "
                               "'26 AUT', '26 SPR' or '26 SUM' — skipped.")
            continue
        r["Session"] = f"{m.group(1)} {m.group(2)}"
        block = r.get("Block", "").strip() or cfg.ALL_BLOCKS
        block = cfg.ALL_BLOCKS if block.lower() == cfg.ALL_BLOCKS.lower() else _number(block)
        if block != cfg.ALL_BLOCKS and block not in cfg.SESSION_BLOCKS[m.group(2)]:
            up.problems.append(f"{where}{who}: block '{r.get('Block')}' isn't a {m.group(2)} block — skipped.")
            continue
        r["Block"] = block
        hrs = _num(r.get("DI hrs/wk", ""))
        if hrs is None or hrs < 0:
            up.problems.append(f"{where}{who}: DI hrs/wk '{r.get('DI hrs/wk', '')}' isn't a number — skipped.")
            continue
        r["DI hrs/wk"] = _number(r["DI hrs/wk"])
        if r.get("Classes"):
            r["Classes"] = _number(r["Classes"])
        if r.get("Discipline"):
            known = _pick(r["Discipline"], programs)
            if known is None:
                up.notes.append(f"{where}{who}: program '{r['Discipline']}' isn't on the list — kept as typed.")
            r["Discipline"] = known or r["Discipline"]
        up.teaching.append(r)
    return up


# --------------------------------------------------------------------------
# what it would change
# --------------------------------------------------------------------------


@dataclass
class Changes:
    new_staff: list[dict] = field(default_factory=list)
    staff_updates: dict[str, dict] = field(default_factory=dict)   # ID -> {column: new value}
    new_teaching: list[dict] = field(default_factory=list)
    teaching_updates: dict[str, dict] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)

    @property
    def count(self) -> int:
        return len(self.new_staff) + len(self.staff_updates) + len(self.new_teaching) + len(self.teaching_updates)


def upload_changes(up: Upload, staff: pd.DataFrame, allocations: pd.DataFrame) -> Changes:
    """What applying `up` would add and change. Not saved."""
    ch = Changes()
    by_emp = {_emp(e): i for i, e in zip(staff["ID"], staff["Employee No"]) if _emp(e)}
    by_key = {name_key(n): i for i, n in zip(staff["ID"], staff["Staff"]) if n}
    rows = staff.set_index("ID")
    name_of: dict[tuple[str, str], str] = {name_key(n): n for n in staff["Staff"] if n}
    staff_cols = [c for c, _ in STAFF_COLS.values()]

    for r in up.staff:
        sid = by_emp.get(r.get("Employee No", "")) or by_key.get(name_key(r["Staff"]))
        if sid is None:
            row = {"ID": wstore.new_id(), "Role": cfg.DEFAULT_ROLE, "FTE": "1", "Active": YES,
                   **{c: r[c] for c in staff_cols if r.get(c)}}
            ch.new_staff.append(row)
            by_key[name_key(row["Staff"])] = row["ID"]
            if row.get("Employee No"):
                by_emp[row["Employee No"]] = row["ID"]
            name_of[name_key(row["Staff"])] = row["Staff"]
            continue
        cur = rows.loc[sid]
        # the name on file stays (it's what teaching and plans are linked by)
        upd = {c: r[c] for c in staff_cols if c != "Staff" and r.get(c) and str(cur[c]) != r[c]}
        if upd:
            ch.staff_updates[sid] = {**ch.staff_updates.get(sid, {}), **upd}

    existing = {(a["Staff"], a["Session"], str(a["Block"]).strip() or cfg.ALL_BLOCKS, a["Subject"].strip().lower()): a
                for a in allocations.to_dict("records")}
    teach_cols = [c for c, _ in TEACHING_COLS.values() if c != "Staff"]
    for r in up.teaching:
        name = name_of.get(name_key(r["Staff"]))
        if name is None:
            ch.problems.append(f"Teaching row {r['_row']}: '{r['Staff']}' isn't on the staff list or the Staff "
                               "sheet — skipped.")
            continue
        key = (name, r["Session"], r["Block"], r.get("Subject", "").strip().lower())
        hit = existing.get(key)
        if hit is None:
            row = {"ID": wstore.new_id(), "Staff": name, **{c: r[c] for c in teach_cols if r.get(c)}}
            ch.new_teaching.append(row)
            existing[key] = row
            continue
        if hit["ID"] in {n["ID"] for n in ch.new_teaching}:
            ch.problems.append(f"Teaching row {r['_row']}: the same person, session, block and subject as an "
                               "earlier row — skipped.")
            continue
        # the subject matched (case aside), so it stays as on file
        upd = {c: r[c] for c in teach_cols if c != "Subject" and r.get(c) and str(hit.get(c, "")) != r[c]}
        if upd:
            ch.teaching_updates[hit["ID"]] = {**ch.teaching_updates.get(hit["ID"], {}), **upd}
    return ch
