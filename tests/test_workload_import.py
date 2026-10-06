"""Tests for the Year plan in the rules and for importing the L&T load
tracking spreadsheet (workload_import), on a made-up workbook."""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

import workload_config as cfg
import workload_import as wi
import workload_rules as rules
import workload_store as wstore
from test_workload import EARLY, adj, alloc, cal, full_year, person, staff
from workload_store import PLANS


def plans(*rows) -> pd.DataFrame:
    base = {"Year": "2026"}
    return pd.DataFrame([{"ID": wstore.new_id(), **base, **r} for r in rows]).reindex(
        columns=PLANS.all_columns, fill_value="")


def with_plan(name, st_, al, plan_row, as_at=EARLY):
    rec = st_[st_["Staff"] == name].iloc[0].to_dict()
    return rules.person_year(rec, al, adj().iloc[0:0], rules.calendar_blocks(cal(), "26"), as_at, plan_row)


# --------------------------------------------------------------------------
# Year plan in the rules
# --------------------------------------------------------------------------


def test_plan_target_overrides_role_and_zero_counts():
    st_ = staff({"Staff": "Ana"})
    p = with_plan("Ana", st_, full_year("Ana", 4), {"Target (h)": "144"})
    assert p["Target"] == 144 and p["Status"] == rules.STATUS_ON
    on_leave = with_plan("Ana", st_, alloc({"Staff": "Ana", "DI hrs/wk": "0"}), {"Target (h)": "0"})
    assert on_leave["Target"] == 0 and on_leave["Status"] == rules.STATUS_ON
    assert with_plan("Ana", st_, full_year("Ana", 16), {"Target (h)": ""})["Target"] == 576


def test_other_allocated_counts_in_total_and_pro_rata_to_date():
    st_ = staff({"Staff": "Ana"})
    c = cal(*[{"Session": "26 AUT", "Block": b, "Start date": d}
              for b, d in [("1", "2026-03-02"), ("2", "2026-03-30")]])
    rec = st_.iloc[0].to_dict()
    p = rules.person_year(rec, alloc({"Staff": "Ana", "DI hrs/wk": "0"}), adj().iloc[0:0],
                          rules.calendar_blocks(c, "26"), pd.Timestamp("2026-03-29"),
                          {"Other allocated (h)": "560"})
    assert p["Other"] == 560 and p["Total"] == 560
    assert p["To date"] == pytest.approx(280)        # 3 of the year's 6 weeks gone
    assert p["Status"] == rules.STATUS_ON           # within 5 % of 576


def test_summarise_year_picks_up_the_years_plan_only():
    st_ = staff({"Staff": "Ana"})
    pl = plans({"Staff": "Ana", "Target (h)": "200"}, {"Staff": "Ana", "Year": "2027", "Target (h)": "10"})
    out = rules.summarise_year(st_, alloc({"Staff": "Ana", "DI hrs/wk": "0"}), adj().iloc[0:0], cal(),
                               "26", EARLY, pl)
    assert out.iloc[0]["Target"] == 200


# --------------------------------------------------------------------------
# names
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw, shown, key",
    [("Smith, Doctor Jane Ann", "Smith, Jane Ann", ("smith", "jane")),
     ("Nguyen, Mr Tom", "Nguyen, Tom", ("nguyen", "tom")),
     ("Lee, Kim", "Lee, Kim", ("lee", "kim")),
     ("Kim Lee", "Kim Lee", ("lee", "kim"))],
)
def test_names(raw, shown, key):
    assert wi.display_name(raw) == shown
    assert wi.name_key(raw) == key


# --------------------------------------------------------------------------
# reading a workbook shaped like the real one
# --------------------------------------------------------------------------


@pytest.fixture
def workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Load Tracking 2026"
    ws["A3"] = "Load tracking - L&T Delivery"
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(["Audit ID", "Employee No", "Name", "Organisation Unit", "Annual load status\n(Jan-Dec 2026)",
               "Required load \n(full year)", "Allocated hours", "Hours underloaded", "Summer hours needed",
               "Articulated plan if underload \n(Yes /No /NA)", "Needs summer teaching?", "Confirmed available?",
               "Now many hours", "Details"])
    ws.append([1, 30000001, "Smith, Doctor Jane", "Science", "On load", 576, 576, 0, None, None, None, None, None, None])
    ws.append([2, 30000002, "Nguyen, Mr Tom", "health science", "Underload", 288, 270, 18, 18, "Yes", None, None,
               None, "Shared summer class"])
    ws.append([None, None, "Lee, Kim", "Business", "Above load", 200, 201, -1, None, "NA", None, None, None, None])
    ws.append([None, None, "Smyth, Jane", "Arts", "On load", 100, 100, 0, None, None, None, None, None, None])
    ws.append([3, 30000003, "Old, Ms Pat", "Arts", None, None, None, 0, None, None, None, None, None, "Resigned"])
    ws.append([None, None, None, None, None, 1164, 1147])                       # totals row
    u = wb.create_sheet("Underload staff - September 26")
    u.append(["Employee No", "Name", "Organisation Unit", "Hours underloaded", "Projected to be met before summer ",
              "Summer teaching needed", "Likelhood of Summer Teaching ", "Subject (if known)"])
    u.append([30000002, "Nguyen, Mr Tom", "Health Science", 18, None, 20, "Confirmed", "EDUC1010"])
    u.append([None, "Lee, Kim", "Business", 1, "Yes", "N/A", None, None])
    u.append([None, "Ghost, Ann", "Arts", 5, None, 5, "Not likily", None])
    path = tmp_path / "load.xlsx"
    wb.save(path)
    return path


def test_parse_reads_both_sheets(workbook):
    r = wi.parse_workbook(workbook)
    by = {p["name"]: p for p in r.people}
    assert set(by) == {"Smith, Jane", "Nguyen, Tom", "Lee, Kim", "Smyth, Jane", "Old, Pat"}
    tom = by["Nguyen, Tom"]
    assert (tom["emp"], tom["unit"], tom["required"], tom["allocated"]) == ("30000002", "Health Science", 288, 270)
    assert (tom["summer"], tom["plan"], tom["likelihood"], tom["subject"]) == (20, "Yes", "Confirmed", "EDUC1010")
    assert by["Lee, Kim"]["plan"] == "N/A" and by["Lee, Kim"]["likelihood"] == "Met before summer"
    assert by["Old, Pat"]["inactive"]


def test_parse_warns_about_unmatched_and_lookalike_names(workbook):
    w = " ".join(wi.parse_workbook(workbook).warnings)
    assert "'Ghost, Ann'" in w and "not imported" in w
    assert "'Smith, Jane' and 'Smyth, Jane' may be the same person" in w


def test_import_adds_staff_and_plans_then_is_idempotent(workbook):
    r = wi.parse_workbook(workbook)
    st_, pl = wstore.empty(wstore.STAFF), wstore.empty(PLANS)
    ch = wi.plan_changes(r, st_, pl, "2026")
    assert len(ch.new_staff) == 5 and len(ch.new_plans) == 5
    tom = next(s for s in ch.new_staff if s["Staff"] == "Nguyen, Tom")
    assert tom["Role"] == cfg.DEFAULT_ROLE and tom["Employee No"] == "30000002"
    assert next(s for s in ch.new_staff if s["Staff"] == "Old, Pat")["Active"] == "FALSE"
    plan = next(p for p in ch.new_plans if p["Staff"] == "Nguyen, Tom")
    assert (plan["Target (h)"], plan["Other allocated (h)"], plan["Summer hrs needed"]) == ("288", "270", "20")

    st2 = pd.DataFrame(ch.new_staff).reindex(columns=wstore.STAFF.all_columns, fill_value="")
    pl2 = pd.DataFrame(ch.new_plans).reindex(columns=PLANS.all_columns, fill_value="")
    again = wi.plan_changes(r, st2, pl2, "2026")
    assert not (again.new_staff or again.new_plans or again.staff_updates or again.plan_updates)


def test_import_matches_existing_staff_and_keeps_their_role(workbook):
    r = wi.parse_workbook(workbook)
    st_ = staff({"Staff": "Tom Nguyen", "Role": "Subject Coordinator", "Home discipline": ""})
    pl = plans({"Staff": "Tom Nguyen", "Target (h)": "300", "Plan": "No"})
    ch = wi.plan_changes(r, st_, pl, "2026")
    assert "Tom Nguyen" not in {s["Staff"] for s in ch.new_staff}
    sid = st_.iloc[0]["ID"]
    assert ch.staff_updates[sid] == {"Employee No": "30000002", "Home discipline": "Health Science"}
    pid = pl.iloc[0]["ID"]
    assert ch.plan_updates[pid]["Target (h)"] == "288" and ch.plan_updates[pid]["Plan"] == "Yes"
