"""Tests for the year grid's Excel export (workload_export)."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import workload_export as wexport


def test_grid_xlsx_keeps_layout_numbers_and_colours():
    grid = pd.DataFrame([
        {"Staff": "Ana", "Role": "Teacher (Ongoing)", "Program": "Science", "AB1": "SCI1001", "AB1 h": "20",
         "AB2": "", "AB2 h": "", "Target": "576", "Total": "1,024", "Variance": "+448", "Status": "over"},
        {"Staff": "Ben", "Role": "Teacher (Casual)", "Program": "Business", "AB1": "BUS1001", "AB1 h": "24",
         "AB2": "BUS1002", "AB2 h": "8", "Target": "", "Total": "192", "Variance": "", "Status": "Casual"},
    ])
    data = wexport.grid_xlsx(grid, ["AB1", "AB2"], over={("Ana", "AB1")}, now="AB2", sheet="Workload 2026")
    ws = load_workbook(BytesIO(data)).active

    assert ws.title == "Workload 2026"
    assert [c.value for c in ws[1]] == ["Name", "Role", "Program", "AB1", "h/wk", "AB2 (now)", "h/wk",
                                        "Target (h)", "Total (h)", "± Target (h)", "Status"]
    assert ws["E2"].value == 20 and ws["I2"].value == 1024 and ws["J2"].value == 448
    assert ws["F2"].value is None and ws["H3"].value is None
    assert ws["E2"].fill.fgColor.rgb.endswith("F8C9A0")             # Ana over in AB1
    assert not ws["E3"].fill.fgColor.rgb.endswith("F8C9A0")         # casual never
    assert ws["F3"].fill.fgColor.rgb.endswith("D5EFD9") and ws["G3"].fill.fgColor.rgb.endswith("D5EFD9")
    assert ws.freeze_panes == "D2"
