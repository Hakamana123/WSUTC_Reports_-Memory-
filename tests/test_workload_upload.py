"""Tests for the bulk staff / teaching upload (workload_upload)."""

from __future__ import annotations

import re
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import workload_config as cfg
import workload_upload as wup
from test_workload import alloc, staff

PROGRAMS = cfg.DISCIPLINES
SESSIONS = ["26 AUT", "26 SPR", "26 SUM"]


def rows(problems: list[str]) -> list[str]:
    """"Staff row 3", … — where each problem is."""
    return [re.match(r"\w+ row \d+", p).group(0) for p in problems]


def filled(staff_rows=(), teaching_rows=()) -> BytesIO:
    """The template with rows written into its Staff and Teaching sheets
    (the example rows left in, as a user might)."""
    wb = load_workbook(BytesIO(wup.template_xlsx(PROGRAMS, SESSIONS)))
    for name, rows in ((wup.STAFF_SHEET, staff_rows), (wup.TEACHING_SHEET, teaching_rows)):
        ws = wb[name]
        for r in rows:
            ws.append(list(r))
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def test_template_has_staff_and_teaching_sheets_with_dropdowns():
    wb = load_workbook(BytesIO(wup.template_xlsx(PROGRAMS, SESSIONS)))
    assert wb.sheetnames[:2] == [wup.STAFF_SHEET, wup.TEACHING_SHEET]
    assert wb["Lists"].sheet_state == "hidden"
    assert [c.value for c in wb[wup.STAFF_SHEET][1]] == list(wup.STAFF_COLS)
    assert [c.value for c in wb[wup.TEACHING_SHEET][1]] == list(wup.TEACHING_COLS)
    assert len(wb[wup.STAFF_SHEET].data_validations.dataValidation) == 3       # role, program, active


def test_example_rows_are_ignored_and_values_cleaned():
    up = wup.parse_upload(filled(
        [("Koh, Josiah", 12345678, "teacher (ongoing)", 0.5, "Theresa", "education studies", "no", "")],
        [("Koh, Josiah", "26 aut", 2, "Business", "BUSM1001", 1, 8, "")],
    ), PROGRAMS)
    assert not up.problems and not up.notes
    s, t = up.staff[0], up.teaching[0]
    assert (s["Staff"], s["Employee No"], s["Role"], s["FTE"], s["Home discipline"], s["Active"]) == \
        ("Koh, Josiah", "12345678", "Teacher (Ongoing)", "0.5", "Education Studies", "FALSE")
    assert (t["Session"], t["Block"], t["Discipline"], t["DI hrs/wk"]) == ("26 AUT", "2", "Business", "8")


def test_bad_rows_are_skipped_with_a_reason():
    up = wup.parse_upload(filled(
        [("Ana", "", "Professor", 1), ("Ben", "", "", 1.5), ("", "", cfg.DEFAULT_ROLE, 1),
         ("Cara", "", "", "", "", "Underwater Basketry")],
        [("Ana", "Autumn", "1", "", "X", "", 4), ("Ana", "26 SUM", "3", "", "X", "", 4),
         ("Ana", "26 AUT", "1", "", "X", "", "lots")],
    ), PROGRAMS)
    assert rows(up.problems) == ["Staff row 3", "Staff row 4", "Staff row 5",
                                                     "Teaching row 3", "Teaching row 4", "Teaching row 5"]
    assert [s["Staff"] for s in up.staff] == ["Cara"]                       # unknown program: kept, noted
    assert "Underwater Basketry" in up.notes[0]


def test_changes_match_existing_staff_and_only_fill_in_whats_given():
    st_ = staff({"Staff": "Koh, Josiah", "Employee No": "111", "FTE": "1", "Supervisor": "Theresa"},
                {"Staff": "Abboud, Antoinette", "Employee No": ""})
    up = wup.parse_upload(filled([
        ("Josiah Koh", 111, "", 0.8, "", ""),                 # by employee no; FTE changes, rest kept
        ("Abboud, Ms Antoinette", 222, "", "", "Chris", ""),  # by name; employee no and supervisor filled in
        ("Lee, Ben", 333, "Teacher (Casual)", "", "", ""),     # new
    ]), PROGRAMS)
    ch = wup.upload_changes(up, st_, alloc().iloc[0:0])
    ids = dict(zip(st_["Staff"], st_["ID"]))
    assert ch.staff_updates == {ids["Koh, Josiah"]: {"FTE": "0.8"},
                                ids["Abboud, Antoinette"]: {"Employee No": "222", "Supervisor": "Chris"}}
    assert len(ch.new_staff) == 1
    new = ch.new_staff[0]
    assert (new["Staff"], new["Role"], new["FTE"], new["Active"]) == ("Lee, Ben", "Teacher (Casual)", "1", "TRUE")


def test_teaching_links_to_staff_and_updates_instead_of_duplicating():
    st_ = staff({"Staff": "Koh, Josiah"})
    al = alloc({"Staff": "Koh, Josiah", "Session": "26 AUT", "Block": "All", "Subject": "EDUC1010", "DI hrs/wk": "12"})
    up = wup.parse_upload(filled(
        [("New, Person", "", "", "", "", "")],
        [("Josiah Koh", "26 AUT", "All", "", "educ1010", "", 10, ""),     # same line: hours updated
         ("Koh, Josiah", "26 AUT", "2", "Business", "BUSM1001", "", 8, ""),
         ("Koh, Josiah", "26 AUT", "2", "Business", "BUSM1001", "", 8, ""),   # repeated in the file
         ("New, Person", "26 SPR", "All", "", "ARTS1001", "", 6, ""),       # someone added by this file
         ("Nobody, At All", "26 SPR", "All", "", "X", "", 6, "")],
    ), PROGRAMS)
    ch = wup.upload_changes(up, st_, al)
    assert ch.teaching_updates == {al["ID"].iloc[0]: {"DI hrs/wk": "10"}}
    assert [(t["Staff"], t["Subject"]) for t in ch.new_teaching] == [("Koh, Josiah", "BUSM1001"),
                                                                     ("New, Person", "ARTS1001")]
    assert rows(ch.problems) == ["Teaching row 5", "Teaching row 7"]
    assert ch.count == 1 + 1 + 2


def test_staff_only_file_is_fine():
    up = wup.parse_upload(filled([("Ana", "", "", "", "", "")]), PROGRAMS)
    assert len(up.staff) == 1 and up.teaching == [] and not up.problems


def test_headers_are_found_by_name():
    df = pd.DataFrame([["Supervisor", "Full name", "Home discipline"], ["Pat", "Ana", "Science"]])
    out = BytesIO()
    with pd.ExcelWriter(out) as xl:
        df.to_excel(xl, sheet_name="staff", header=False, index=False)
    out.seek(0)
    up = wup.parse_upload(out, PROGRAMS)
    assert up.staff[0]["Supervisor"] == "Pat" and up.staff[0]["Home discipline"] == "Science"
