"""The dashboard's year grid as an Excel workbook, laid out as on screen:
Name / Role / Program, then a subject column and an h/wk column per block in
calendar order, then target, total, ± and status. Orange = a block above the
role's weekly limit; green = the current block. Names and the header row are
frozen. Pure: a DataFrame in, bytes out.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ORANGE = PatternFill("solid", fgColor="F8C9A0")
GREEN = PatternFill("solid", fgColor="D5EFD9")
HEADER = PatternFill("solid", fgColor="EDEDED")

TITLES = {"Staff": "Name", "Target": "Target (h)", "Total": "Total (h)", "Variance": "± Target (h)"}


def _value(v):
    """Numbers shown as text on screen ("16", "1,024", "+12") back to numbers."""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, str):
        try:
            return float(v.replace(",", "")) if v.strip() else None
        except ValueError:
            return v
    return v


def grid_xlsx(grid: pd.DataFrame, labels: list[str], over: set[tuple[str, str]], now: str | None,
              sheet: str = "Workload") -> bytes:
    """`grid` as the dashboard shows it (columns Staff, Role, Program, then
    "<label>" / "<label> h" per block, then Target, Total, Variance, Status)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    cols = list(grid.columns)
    hours = {f"{lb} h": lb for lb in labels}

    for j, c in enumerate(cols, 1):
        title = "h/wk" if c in hours else TITLES.get(c, c)
        if c == now:
            title += " (now)"
        cell = ws.cell(row=1, column=j, value=title)
        cell.font, cell.fill = Font(bold=True), HEADER
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(grid.itertuples(index=False), 2):
        name = row[0]
        for j, (c, v) in enumerate(zip(cols, row), 1):
            cell = ws.cell(row=i, column=j, value=_value(v))
            block = hours.get(c, c if c in labels else None)
            if block and block == now:
                cell.fill = GREEN
            if c in hours and (name, hours[c]) in over:
                cell.fill, cell.font = ORANGE, Font(bold=True)

    for j, c in enumerate(cols, 1):
        width = 7 if c in hours else 28 if c in labels or c in ("Staff", "Program") else 14
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "D2"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
